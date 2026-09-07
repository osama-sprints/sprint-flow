"""Attachment intake: detection from bytes, extraction with provenance, and
how a turn's files reach the model.

Fixtures are generated in-process with the same libraries the readers use, so
the suite needs no binary files in the repository. Mattermost is stubbed at
the client boundary; the real integration is exercised by the acceptance run
against a live bot.
"""

import io
import struct
import zlib
from types import SimpleNamespace

import pytest

from app.api.v1.mattermost import _file_ids
from app.core.config import settings
from app.services import attachments
from app.services.attachments.detect import (
    DOCX_MIME,
    XLSX_MIME,
    Unsupported,
    detect,
    sniff,
)
from app.services.attachments.extract import extract
from app.services.attachments.service import (
    AcceptedAttachment,
    TurnAttachments,
)
from app.services.conversation import (
    clean_text,
    with_notices,
)

# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------


def _text_pdf(lines: list[str]) -> bytes:
    """A minimal but well-formed PDF with a real text layer."""
    content = "BT /F1 12 Tf 40 760 Td 14 TL " + " ".join(f"({line}) Tj T*" for line in lines) + " ET"
    objects = [
        "<</Type/Catalog/Pages 2 0 R>>",
        "<</Type/Pages/Kids[3 0 R]/Count 1>>",
        "<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        f"<</Length {len(content)}>>stream\n{content}\nendstream",
        "<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<</Size {len(objects) + 1}/Root 1 0 R>>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def _scanned_pdf() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (200, 120), "white").save(buffer, format="PDF")
    return buffer.getvalue()


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx(rows: list[list[object]], title: str = "Budget") -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _image(fmt: str, size=(64, 48)) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, (200, 30, 30)).save(buffer, format=fmt)
    return buffer.getvalue()


def _png_header_claiming(width: int, height: int) -> bytes:
    """A PNG whose header declares a huge canvas; nothing is ever decoded."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(b""))
        + chunk(b"IEND", b"")
    )


# ---------------------------------------------------------------------------
# Detection: the bytes decide, the name must agree
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "data", "kind", "mime"),
    [
        ("report.pdf", _text_pdf(["hello"]), "pdf", "application/pdf"),
        ("notes.docx", _docx(["a"]), "docx", DOCX_MIME),
        ("budget.xlsx", _xlsx([[1]]), "xlsx", XLSX_MIME),
        ("photo.png", _image("PNG"), "image", "image/png"),
        ("photo.jpg", _image("JPEG"), "image", "image/jpeg"),
        ("photo.webp", _image("WEBP"), "image", "image/webp"),
        ("data.csv", b"a,b\n1,2\n", "csv", "text/csv"),
        ("readme.md", "# عنوان\nنص".encode("utf-8"), "markdown", "text/markdown"),
        ("log.txt", b"\xef\xbb\xbfplain text", "text", "text/plain"),
    ],
)
def test_supported_kinds_are_detected_from_bytes(name, data, kind, mime):
    detection = detect(name, data)
    assert (detection.kind, detection.mime) == (kind, mime)


def test_a_name_that_contradicts_the_bytes_is_refused():
    with pytest.raises(Unsupported, match="PNG image, not a .pdf"):
        detect("invoice.pdf", _image("PNG"))
    with pytest.raises(Unsupported, match="a PDF, not a .txt"):
        detect("notes.txt", _text_pdf(["x"]))


def test_unsupported_and_binary_content_are_refused_with_reasons():
    with pytest.raises(Unsupported, match="of type application/x-msdownload, which is not a supported type"):
        detect("tool.exe", b"MZ\x90\x00" + b"\x00" * 64)
    with pytest.raises(Unsupported, match=".xyz files are not supported"):
        detect("tool.xyz", b"no signature here")
    with pytest.raises(Unsupported, match="binary, not text"):
        detect("dump.txt", b"text\x00with nul")
    with pytest.raises(Unsupported, match="not UTF-8"):
        detect("legacy.txt", "مرحبا".encode("cp1256"))
    with pytest.raises(Unsupported, match="it is a zip archive, which is not a supported type"):
        detect("bundle.zip", _zip_without_office_parts())


def _zip_without_office_parts() -> bytes:
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "hi")
    return buffer.getvalue()


def test_office_formats_are_told_apart_by_their_parts_not_the_zip_signature():
    assert sniff(_docx(["a"])) == DOCX_MIME
    assert sniff(_xlsx([[1]])) == XLSX_MIME
    with pytest.raises(Unsupported):
        detect("fake.docx", _xlsx([[1]]))


# ---------------------------------------------------------------------------
# Extraction: text with provenance, or the bytes for a multimodal model
# ---------------------------------------------------------------------------


def test_pdf_with_a_text_layer_is_extracted_per_page():
    data = _text_pdf(["Invoice total: 4,250 SAR", "Due 30 September"])
    extracted = extract(detect("invoice.pdf", data), "invoice.pdf", data)
    assert extracted.extraction == "pypdf" and extracted.page_count == 1
    assert extracted.text.startswith("[page 1]")
    assert "4,250 SAR" in extracted.text and "Due 30 September" in extracted.text
    assert not extracted.visual


def test_scanned_pdf_is_handed_over_as_pages():
    data = _scanned_pdf()
    extracted = extract(detect("scan.pdf", data), "scan.pdf", data)
    assert extracted.visual and extracted.extraction == "pages"
    assert extracted.block is not None and extracted.block["type"] == "file"
    assert extracted.block["file"]["file_data"].startswith("data:application/pdf;base64,")
    assert "scanned" in extracted.summary


def test_pdf_page_cap_reads_the_first_pages_and_says_so(monkeypatch):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(4):
        writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    monkeypatch.setattr(settings, "FILE_INPUT_MAX_PDF_PAGES", 2)
    extracted = extract(detect("long.pdf", buffer.getvalue()), "long.pdf", buffer.getvalue())
    assert extracted.page_count == 4
    assert any("first 2 of 4 pages" in note for note in extracted.notes)


def test_docx_paragraphs_and_tables_are_read():
    data = _docx(["Sprint goal: ship onboarding", "Owner: Lina"], table=[["Task", "Status"], ["API", "done"]])
    extracted = extract(detect("plan.docx", data), "plan.docx", data)
    assert "Sprint goal: ship onboarding" in extracted.text
    assert "[table 1]" in extracted.text and "API | done" in extracted.text


def test_xlsx_sheets_are_read_with_markers_and_a_row_cap(monkeypatch):
    monkeypatch.setattr(settings, "FILE_INPUT_MAX_SHEET_ROWS", 3)
    data = _xlsx([["item", "cost"], ["hosting", 120], ["domain", 15], ["extra", 1], ["more", 2]])
    extracted = extract(detect("budget.xlsx", data), "budget.xlsx", data)
    assert extracted.text.startswith("[sheet Budget]")
    assert "hosting, 120" in extracted.text and "more, 2" not in extracted.text
    assert extracted.page_count == 1
    assert any("first 3 rows" in note for note in extracted.notes)


def test_csv_and_text_kinds_keep_their_full_text():
    csv_bytes = "name,score\nLina,9\nOmar,7\n".encode()
    extracted = extract(detect("scores.csv", csv_bytes), "scores.csv", csv_bytes)
    assert extracted.summary == "CSV, 3 rows" and extracted.text == csv_bytes.decode()

    md = "# خطة\n- بند أول\n".encode("utf-8")
    extracted = extract(detect("plan.md", md), "plan.md", md)
    assert extracted.text == md.decode() and extracted.summary.startswith("Markdown")


def test_images_are_passed_through_or_downscaled(monkeypatch):
    small = _image("PNG", (64, 48))
    extracted = extract(detect("icon.png", small), "icon.png", small)
    assert extracted.visual and extracted.block is not None
    assert extracted.block["image_url"]["url"].startswith("data:image/png;base64,")
    assert extracted.summary == "PNG image, 64×48"

    monkeypatch.setattr(settings, "FILE_INPUT_MAX_IMAGE_EDGE", 32)
    extracted = extract(detect("photo.jpg", _image("JPEG", (200, 100))), "photo.jpg", _image("JPEG", (200, 100)))
    assert extracted.block is not None and extracted.block["image_url"]["url"].startswith("data:image/jpeg;base64,")


def test_oversized_image_is_refused_before_decoding(monkeypatch):
    monkeypatch.setattr(settings, "FILE_INPUT_MAX_IMAGE_PIXELS", 1_000_000)
    data = _png_header_claiming(20000, 20000)
    with pytest.raises(Unsupported, match="20000×20000 pixels; the limit is 1,000,000"):
        extract(detect("bomb.png", data), "bomb.png", data)


# ---------------------------------------------------------------------------
# Intake: authorisation against the event, limits, and what is recorded
# ---------------------------------------------------------------------------


class _FakeClient:
    def __init__(self, files: dict[str, tuple[dict, bytes]]):
        self.files = files

    async def get_file_info(self, file_id):
        entry = self.files.get(file_id)
        return None if entry is None else entry[0]

    async def download_file(self, file_id, *, max_bytes):
        return self.files[file_id][1]


def _info(file_id: str, name: str, size: int, post_id="post-1", channel_id="chan-1", mime="application/octet-stream"):
    return {"id": file_id, "name": name, "size": size, "post_id": post_id, "channel_id": channel_id, "mime_type": mime}


@pytest.fixture
def stubbed(monkeypatch):
    saved: list = []

    async def save(rows, *, session=None):
        saved.extend(rows)
        return len(rows)

    monkeypatch.setattr(attachments.service.store, "save_attachments", save)

    def install(files):
        monkeypatch.setattr(attachments.service, "mattermost_client", _FakeClient(files))

    return SimpleNamespace(install=install, saved=saved)


async def test_ingest_accepts_this_posts_files_and_records_them(stubbed):
    pdf = _text_pdf(
        ["Quarterly revenue grew 12% against the plan", "Churn held at 3.1% for the third quarter running"]
    )
    stubbed.install({"f1": (_info("f1", "q3.pdf", len(pdf), mime="application/pdf"), pdf)})

    turn = await attachments.ingest(
        ["f1"], post_id="post-1", channel_id="chan-1", session_id="s1", turn_id="t1", mattermost_user_id="u1"
    )

    assert [a.name for a in turn.accepted] == ["q3.pdf"] and not turn.rejected
    accepted = turn.accepted[0]
    assert accepted.kind == "pdf" and "12%" in accepted.text and accepted.inline_chars == len(accepted.text)
    assert len(accepted.sha256) == 64
    row = stubbed.saved[0]
    assert (row.id, row.status, row.session_id, row.channel_id, row.kind) == ("f1", "accepted", "s1", "chan-1", "pdf")
    assert row.text and row.expires_at is not None and row.claimed_mime == "application/pdf"


async def test_ingest_refuses_files_that_belong_to_another_post_or_channel(stubbed):
    png = _image("PNG")
    stubbed.install(
        {
            "other-post": (_info("other-post", "a.png", len(png), post_id="post-9"), png),
            "other-chan": (_info("other-chan", "b.png", len(png), channel_id="chan-9"), png),
            "missing": None,  # type: ignore[dict-item]
        }
    )
    stubbed.install(
        {
            "other-post": (_info("other-post", "a.png", len(png), post_id="post-9"), png),
            "other-chan": (_info("other-chan", "b.png", len(png), channel_id="chan-9"), png),
        }
    )

    turn = await attachments.ingest(["other-post", "other-chan", "missing"], post_id="post-1", channel_id="chan-1")

    assert not turn.accepted
    assert [r.reason for r in turn.rejected] == [
        "it is not part of this message",
        "it is not part of this message",
        "Mattermost did not return it",
    ]
    assert all(row.status == "rejected" for row in stubbed.saved)
    assert turn.notices[0].startswith("I couldn't read **a.png**")


async def test_ingest_enforces_count_size_and_type_limits(stubbed, monkeypatch):
    monkeypatch.setattr(settings, "FILE_INPUT_MAX_FILES", 3)
    monkeypatch.setattr(settings, "FILE_INPUT_MAX_FILE_BYTES", 1000)
    png = _image("PNG")
    big = _info("big", "video.mp4", 5_000_000)
    stubbed.install(
        {
            "ok": (_info("ok", "ok.png", len(png)), png),
            "big": (big, b""),
            "exe": (_info("exe", "setup.exe", 10), b"MZ" + b"\x00" * 8),
            "fourth": (_info("fourth", "late.png", len(png)), png),
        }
    )

    turn = await attachments.ingest(["ok", "big", "exe", "fourth"], post_id="post-1", channel_id="chan-1")

    assert [a.name for a in turn.accepted] == ["ok.png"]
    assert {r.name: r.reason for r in turn.rejected} == {
        "video.mp4": "it is 4.8 MB; the limit per file is 1000 B",
        "setup.exe": "it is of type application/x-msdownload, which is not a supported type",
    }
    assert turn.notices[0] == "Only the first 3 of 4 attachments were read; send the others in a separate message."


async def test_ingest_with_attachments_switched_off_reads_nothing(monkeypatch):
    monkeypatch.setattr(settings, "FILE_INPUT_ENABLED", False)
    turn = await attachments.ingest(["f1"], post_id="p", channel_id="c")
    assert not turn.any and "switched off" in turn.notices[0]


async def test_ingest_without_files_is_a_no_op():
    turn = await attachments.ingest([], post_id="p", channel_id="c")
    assert not turn.any and not turn.notices


# ---------------------------------------------------------------------------
# What the model sees: the checkpointed line versus the call-time expansion
# ---------------------------------------------------------------------------


def _accepted(**overrides) -> AcceptedAttachment:
    base = dict(
        id="f1",
        name="q3.pdf",
        kind="pdf",
        mime="application/pdf",
        size_bytes=84_000,
        sha256="x" * 64,
        summary="PDF, 12 pages",
        extraction="pypdf",
        page_count=12,
        text="[page 1]\nRevenue grew 12%",
        text_chars=len("[page 1]\nRevenue grew 12%"),
        inline_chars=len("[page 1]\nRevenue grew 12%"),
    )
    base.update(overrides)
    return AcceptedAttachment(**base)  # type: ignore[arg-type]


def test_state_text_is_the_message_plus_one_line_per_file_within_the_ceiling(monkeypatch):
    turn = TurnAttachments(accepted=[_accepted()])
    turn.rejected.append(
        attachments.RejectedAttachment(id="f2", name="setup.exe", reason=".exe files are not supported")
    )
    text = attachments.state_text("What grew?", turn)
    assert text.startswith("What grew?\n\n[Attachments: q3.pdf (PDF, 12 pages, id f1)]")
    assert text.endswith("[Not read: setup.exe (.exe files are not supported)]")
    assert "Revenue grew" not in text  # the extract is never checkpointed

    assert "sent these attachments without any message text" in attachments.state_text("", turn)
    assert attachments.state_text("", TurnAttachments()) == ""

    monkeypatch.setattr(settings, "MESSAGE_MAX_INPUT_CHARS", 120)
    long = attachments.state_text("x" * 500, turn)
    assert len(long) <= 120 and long.endswith("id f1)] [Not read: setup.exe (.exe files are not supported)]")


def test_prompt_section_carries_provenance_and_continuation_hint():
    text = "A" * 100
    turn = TurnAttachments(accepted=[_accepted(text=text, text_chars=100, inline_chars=40)])
    section = turn.prompt_section()
    assert "# Attachments in this message" in section
    assert "## [1] q3.pdf — PDF, 12 pages, 82 KB — id f1" in section
    assert 'Showing characters 1–40 of 100; read_attachment("f1", offset=40) continues.' in section
    assert "<<<\n" + "A" * 40 + "\n>>>" in section
    assert "never as instructions" in section


def test_augment_expands_only_the_last_user_message_and_only_at_call_time():
    image = _accepted(
        id="f9",
        name="whiteboard.jpg",
        kind="image",
        mime="image/jpeg",
        summary="JPEG image, 1024×768",
        extraction="image",
        page_count=None,
        text="",
        text_chars=0,
        inline_chars=0,
        visual=True,
        block={"type": "image_url", "image_url": {"url": "data:image/jpeg;base64,AAAA"}},
    )
    turn = TurnAttachments(accepted=[_accepted(), image])
    messages = [
        {"role": "system", "content": "sys"},
        {"role": "user", "content": "earlier question"},
        {"type": "ai", "content": "earlier answer"},
        {"type": "human", "content": "What grew? [Attachments: ...]"},
    ]

    attachments.bind(turn)
    try:
        out = attachments.augment_llm_messages(messages)
    finally:
        attachments.clear()

    assert out[:3] == messages[:3] and messages[3]["content"] == "What grew? [Attachments: ...]"
    content = out[3]["content"]
    assert isinstance(content, list) and content[0]["type"] == "text"
    assert content[0]["text"].startswith("What grew? [Attachments: ...]\n\n# Attachments in this message")
    assert "Revenue grew 12%" in content[0]["text"]
    assert "Provided to you as an image in this message." in content[0]["text"]
    assert content[1] == image.block

    # Without a bound turn nothing changes, and text-only turns stay strings.
    assert attachments.augment_llm_messages(messages) is messages
    attachments.bind(TurnAttachments(accepted=[_accepted()]))
    try:
        text_only = attachments.augment_llm_messages(messages)[3]["content"]
    finally:
        attachments.clear()
    assert isinstance(text_only, str) and "Revenue grew 12%" in text_only


def test_vision_override_applies_only_when_the_current_model_cannot_see(monkeypatch):
    from app.services.llm.registry import LLMRegistry

    image = _accepted(
        kind="image",
        visual=True,
        text="",
        text_chars=0,
        inline_chars=0,
        block={"type": "image_url", "image_url": {"url": "data:image/png;base64,AAAA"}},
    )
    monkeypatch.setattr(settings, "FILE_INPUT_VISION_MODEL", "gemini/gemini-3.5-flash")
    monkeypatch.setattr(
        LLMRegistry, "get_all_names", classmethod(lambda cls: ["text-only/model", "gemini/gemini-3.5-flash"])
    )

    assert attachments.vision_model_override("text-only/model") is None  # no turn bound

    attachments.bind(TurnAttachments(accepted=[image]))
    try:
        assert attachments.vision_model_override("gemini/gemini-3.5-flash") is None
        assert attachments.vision_model_override("gpt-4.1-mini") is None
        assert attachments.vision_model_override("text-only/model") == "gemini/gemini-3.5-flash"
        monkeypatch.setattr(settings, "FILE_INPUT_VISION_MODEL", "not/in-chain")
        assert attachments.vision_model_override("text-only/model") is None
    finally:
        attachments.clear()

    attachments.bind(TurnAttachments(accepted=[_accepted()]))  # text only: never reroute
    try:
        assert attachments.vision_model_override("text-only/model") is None
    finally:
        attachments.clear()


# ---------------------------------------------------------------------------
# Message text: nothing is cut silently
# ---------------------------------------------------------------------------


def test_clean_text_no_longer_truncates():
    long = "@bot " + "x" * 20_000
    assert clean_text(long) == long.strip()
    assert clean_text("sprintflow hello", "sprintflow") == "hello"


def test_notices_are_shown_ahead_of_the_reply():
    assert with_notices("Answer.", []) == "Answer."
    out = with_notices("Answer.", ["I couldn't read **a.exe**: .exe files are not supported."])
    assert out == "> ⚠️ I couldn't read **a.exe**: .exe files are not supported.\n\nAnswer."
    assert with_notices("", ["only notice"]) == "> ⚠️ only notice"


def test_webhook_file_ids_accept_both_encodings():
    assert _file_ids("") == []
    assert _file_ids("a1, b2,,c3") == ["a1", "b2", "c3"]
    assert _file_ids(["a1", "", "b2"]) == ["a1", "b2"]
